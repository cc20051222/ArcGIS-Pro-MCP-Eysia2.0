"""Author the requested two manuals using the managed document runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'配套交付/ArcGIS-Pro-MCP-2.0-Skill与面板-0.7.1候选/05_说明'
OUT.mkdir(parents=True,exist_ok=True)

def new(title,subtitle):
    d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.78)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=d.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(4)
    d.styles['Normal'].paragraph_format.line_spacing=1.02
    d.styles['Heading 1'].paragraph_format.space_before=Pt(14);d.styles['Heading 1'].paragraph_format.space_after=Pt(7)
    d.styles['Heading 2'].paragraph_format.space_before=Pt(9);d.styles['Heading 2'].paragraph_format.space_after=Pt(4)
    s.header_distance=Inches(.25);s.footer_distance=Inches(.25)
    d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(18);d.styles['Heading 2'].font.size=Pt(13)
    for style in d.styles:
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):style.element.pPr.remove(border)
    s.header.paragraphs[0].text='ArcGIS Pro MCP 2.0  /  Skill 与任务面板'
    s.header.paragraphs[0].runs[0].font.size=Pt(9)
    foot=s.footer.paragraphs[0];foot.alignment=2;foot.add_run('0.7.1 配套候选  ·  ')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
    d.add_heading(title,0);d.add_paragraph(subtitle,'Subtitle')
    d.add_paragraph('2026年10月19日｜适用对象：本项目 2.0 用户与部署维护人员')
    return d

def h(d,title,body):
    d.add_heading(title,2)
    for p in body.split('\n\n'):d.add_paragraph(p)

def page(d,title):
    d.add_heading(title,1)

def cmd(d,text):
    p=d.add_paragraph();r=p.add_run(text);r.font.name='Consolas';r.font.size=Pt(9)
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(9)

def matrix(d,headers,rows,widths):
    t=d.add_table(rows=1, cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=txt
    for i,row in enumerate(t.rows):
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            c.vertical_alignment=1;pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if i==0 else 'F5F7FA' if i%2==0 else 'FFFFFF');pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.font.size=Pt(10);r.bold=i==0
    d.add_paragraph()

def picture(d,name,caption):
    p=ROOT/'.runtime/progress-panel/intake-070'/name
    if p.exists():
        page(d,'实际界面与阅读方法')
        d.add_picture(str(p),width=Inches(6.5));d.add_paragraph(caption)
        h(d,'先看业务链，再看检查与交付','左侧保留全部业务步骤及依赖，右侧按需查看步骤详情、检查与修复、内部进度和资料与交付。顶部当前执行信息与行内高亮都依据真实观测，查看其它行只改变详情选择。\n\n“检查与修复”用于区分首次检查、失败、修复及新证据复检；“资料与交付”用于查看需求、数据路径、补充建议、工具选择以及实际成果位置。内部日志属于辅助解释，不能把工具成功直接认定为成果通过。')

d=new('插件 Skill 与任务面板安装及使用教程','从安装、资料确认到复检交付的完整操作指南')
h(d,'先认清本次交付','本目录将原版 r29 安装包、0.7.1 Skill、独立面板、新插件构件与本教程放在一起。原版 r29 为原样副本；新增构件是明确标识的配套候选，不覆盖封存 Release，也不冒称已正式发布的新 r29。原版安装不包含本次面板改动。')
matrix(d,['内容','位置','作用'],[
 ('原版插件安装包','01_原版r29','保留原版安装、连接与兼容流程'),
 ('Skill 0.7.1','02_Skill','业务规划、问询、检查与对话展示'),
 ('独立面板','03_Panel','可选的独立只读观察窗口'),
 ('新插件构件','04_Plugin','Pro 内任务面板与实际事件'),
 ('两份说明','05_说明','安装教程、功能特色和实现原理')],[1.5,1.5,3.8])
h(d,'安装条件','本次新插件只针对本机验证过的 ArcGIS Pro 3.5 / net8；其它 Pro 版本使用原版兼容流程，不强装新构件。独立面板需要 .NET 8 Windows Desktop Runtime；Skill 发布脚本需要本地 Python 3.10 或以上版本。安装脚本不下载依赖。\n\nAI 客户端必须能读取项目 Skill，并具备调用本机 MCP 的能力。只有远程 MCP 工具而没有本地文件/脚本权限的客户端，无法直接运行当前任务发布脚本；此时必须明确说明面板同步未具备。')

page(d,'1  先核包，再安装三个部件')
h(d,'安装前准备','将整个配套目录保存到 D 盘，保留原目录结构；先保存工程、正常退出 Pro，脚本确认进程 remaining=0。选定 D 盘运行目录与 AI 工作区。输入资料默认只读保留，安装不代表授权任何 GIS 数据改写。')
h(d,'第一步：原版安装与连接','新电脑先解压 01_原版r29 中的 ZIP，按其中原教程安装原版插件并配置 MCP 连接。已正常使用本项目插件的电脑无需重复部署原版。MCP 保持本机 127.0.0.1:6520/mcp；不开放公网，不增加任意 Python 通道。')
h(d,'第二步：检查配套安装计划','在配套目录中打开 PowerShell。示例路径须换为自己的 D 盘位置；Plan 只核对包内哈希并展示计划，不安装。脚本路径带空格或中文时必须保留引号。')
cmd(d,"& '.\\安装配套.ps1' -Action Plan `\n  -WorkRoot 'D:\\MCP运行记录' -AgentWorkspace 'D:\\你的工作区'")
h(d,'第三步：安装插件、Skill 和面板','确认保存并关闭 Pro 后运行下方命令。AllowOnlyThisAddinOnC 明确允许仅本插件 GUID 的 AddIns 与 AssemblyCache 两个目录更新；项目、备份、日志与临时文件在 D 盘。UpdateExisting 表示先备份后更新不同版本；不带此参数时拒绝覆盖不同文件。')
cmd(d,"& '.\\安装配套.ps1' -Action InstallAll `\n  -WorkRoot 'D:\\MCP运行记录' -AgentWorkspace 'D:\\你的工作区' `\n  -UpdateExisting -AllowOnlyThisAddinOnC")
h(d,'安装结果在哪里','插件：当前用户 Documents/ArcGIS/AddIns/ArcGISPro/{BAA5628C-3C08-4AD5-A6C0-915ADA475709}/ArcGISProMCP.Compatibility.esriAddInX；缓存仅对应 GUID 目录。Skill：所选工作区/.agents/skills/arcgis-pro-2-workflows；独立面板：运行目录/panel。每次安装在运行目录/transactions 中保存带时间与唯一编号的 receipt.json 及备份。\n\n也可分开使用 InstallSkill、InstallPanel、InstallPlugin。安装文件成功不代表客户端已加载 Skill，也不代表新 GIS 任务已验收。不要在本次配套之后再执行原版部署，以免降回旧插件。')

page(d,'2  启动 Pro、Skill 与持续观察面板')
h(d,'统一任务与事件路径','用配套启动器启动已存在的 D 盘工程，使本次 Pro 进程继承 D 盘日志、事件、任务记录、作业、审计和 scratch 设置。已经打开的 Pro 无法追补进程环境变量，应先保存并正常退出。启动器不改真实客户端配置。')
cmd(d,"& '.\\启动配套.ps1' -Mode Pro -WorkRoot 'D:\\MCP运行记录' `\n  -Project 'D:\\工程\\你的工程.aprx'")
h(d,'在客户端启用 Skill','让客户端打开同一个 AgentWorkspace，并重新加载技能列表。项目级 Skill 的入口为 .agents/skills/arcgis-pro-2-workflows/SKILL.md。本项目 Codex 已验证该项目级位置；其他客户端按其实际技能发现机制导入，不保证仅复制文件就自动启用。随后在 Pro 的 MCP 页签启动服务器，检查工具发现并打开“任务面板”。')
h(d,'独立面板是可选观察入口','Pro 内面板已经可以观看；需要单独窗口时再运行以下命令。两个窗口都只读观察同一文件，不额外开启服务，不执行 GIS 或取消任务。')
cmd(d,"& '.\\启动配套.ps1' -Mode Panel -WorkRoot 'D:\\MCP运行记录'")
h(d,'Agent 发布规则','Agent 将同一原始记录交给 Skill/scripts/publish_workflow.py，并发布到 D:\\MCP运行记录\\task-record.json；脚本原子替换并保留历史。事件目录为 D:\\MCP运行记录\\events。聊天阶段卡由 Agent 在关键节点直接发送，面板不会修改既有聊天消息。')
cmd(d,"python -B '<Skill路径>\\scripts\\publish_workflow.py' `\n  'D:\\MCP运行记录\\本次任务.json' `\n  --target 'D:\\MCP运行记录\\task-record.json'")
h(d,'首次验证','确认 initialize 声明进度能力、tools/list 能发现当前工具；再发出已知安全的只读查询。面板有当前任务、真实时间与正确输出根，聊天显示相同业务步骤，才算完成连接与显示检查。不要把工具总数等同所有功能都可执行。')

page(d,'3  用户怎样完成一次业务任务')
h(d,'开始：需求、数据、成果位置','Agent 集中问“要做什么、数据在哪里、结果放哪里”。用户已经提供的内容直接复述沿用，只问未知或变化项。例：使用 D 盘某 DEM 和研究区边界，分析指定出口的汇水范围，结果保存到指定文件夹。未明确的坐标系、出口、时相、方法或单位按实际需要补问。')
h(d,'资料不全：先给可操作清单','缺失项应列数据名称、用途、必需/可选、推荐格式/范围/分辨率或精度、影响步骤和可接受替代。资料欠缺时仍可做独立只读预检，相关处理步骤保持等待。不得自动找一份示例数据冒充用户资料。')
h(d,'制定方案：现有工具与科学方法优先','先检查实际 schema、许可、实现与数据适用性，再选择处理链。缺少工具时先评估现有受控工具组合，说明等效条件、差异及质量影响；确无替代才提出开发建议，并取得用户明确开发同意。开发获准后仍需实现与检验，不能直接标成工具已支持。')
h(d,'观看：当前步骤、检查与修复','运行行有绿色底色、左侧色条及“▶ 当前步骤”，顶部同时显示当前步骤；多步真实并发可同时高亮。点击其它行查看细节不会改变其执行状态；“跟随当前步骤”控制是否自动定位。失败保留检查证据，修复后下游失效项需要新证据复检。')
picture(d,'native-running-live300.png','真实 Pro 截图：有界只读事件压力检查进行中；该测试不代表新增业务范围或水文算法已执行。')

page(d,'4  交付、排障与可恢复更新')
h(d,'结束后逐件告知位置','聊天和“资料与交付”页都应列成果名称、完整绝对路径、格式、核验状态及未完成成员。计划路径只说明将写到哪里，不表示文件已生成；正式图件需要实际视觉检查，检查报告需要真实重新读取。质量自检通过形成 PASS CANDIDATE，外部验收另列。')
matrix(d,['现象','先检查什么'],[
 ('有插件但没有任务面板','确认运行构件为本次候选；原版 r29 没有本次界面'),
 ('面板停在旧任务','核对 Pro 与发布器的 task-record 路径；按准确 taskId 显式切换任务'),
 ('工具执行但业务状态不更新','Agent 是否及时发布、requestIds 是否准确绑定；500ms 读取不能替代发布'),
 ('显示旧记录未确认','补全 intake 的需求、数据和输出根；不修改历史为虚假通过'),
 ('安装拒绝或文件被占用','保存并正常关闭 Pro；查事务记录，不强杀或盲重试'),
 ('资料或工具欠缺','补齐资料或寻找科学等效替代；新工具开发先获明确同意')],[2.0,4.8])
h(d,'回滚与卸除配套','保留安装时的 receipt.json 与备份。保存并正常关闭 Pro 后，按下例恢复已登记文件。脚本先核安装后哈希；用户后来改过文件会停止，避免覆盖。未知文件不删除，空目录可以保留。部分失败的事务先根据记录对账，不能保证无需判断就直接恢复。')
cmd(d,"& '.\\安装配套.ps1' -Action Restore -WorkRoot 'D:\\MCP运行记录' `\n  -Receipt 'D:\\MCP运行记录\\transactions\\<事务>\\receipt.json' `\n  -AllowOnlyThisAddinOnC")
h(d,'版本与验证边界','0.7.1 本机 net6/net8 编译及独立面板编译均为零警告零错误；97 项离线 Skill 辅助检查、16 项进度协议检查通过；真实 Pro 239 工具可发现，既有三区 12 次只读首检/复检，15 个源成员哈希未变，并取得真实运行高亮截帧。发布仍为候选；不把这些证据推广为所有客户端、所有Pro版本或所有业务算法均通过，综合实验的技术复检与独立验收仍需分别记录。')
d.save(OUT/'01_插件Skill与面板_安装使用教程.docx')

d=new('Skill 与任务面板\n功能特色与实现原理','让业务流程、真实执行与成果检查保持一致')
h(d,'产品定位','Skill 将用户业务目标转成当前工具能力能够执行的 GIS 流程；面板持续展示这份流程、真实工具观测、检查结果与修复复检。用户主要关心“做到了哪一步、依据是什么、哪里有问题、结果在哪里”，不必靠阅读后台工具日志推测进展。')
h(d,'本次六项改进','一是执行步骤与选中查看行独立，真实 RUNNING 状态有持续高亮；二是按需求、数据、适用依据和当前工具制定计划，替代优先；三是任务开始确认成果目录，结束逐件列准确文件位置；四是任务开始确认需求与数据，欠缺时给出补充清单；五是面板与聊天共用任务合同；六是将原版安装与新版配套整理在同一交付目录，并提供可核对的安装记录。')
matrix(d,['部件','负责内容','不代表什么'],[
 ('Skill','问询、方法选择、调用约束、检查规则、聊天更新','不是后台算法引擎，也不自动授予写权限'),
 ('任务发布器','校验记录、推导展示事实、原子发布、历史保留','不是 GIS 执行器或独立认证方'),
 ('Pro 原生面板','业务步骤、质量闭环、资料交付、请求实况','不是自动补数、自动修复或聊天输入通道'),
 ('独立面板','只读相同 D 盘记录与事件','不另开网络端口，不自行启动 GIS'),
 ('既有 MCP 工具','受控操作及真实结果','已注册数量不等同所有算法均已实现')],[1.25,2.6,2.95])

page(d,'1  从业务目标形成可执行步骤')
h(d,'适用场景','可以用于资料盘点、矢量处理、已支持的栅格/统计入口、质量检查、制图与交付等任务。不同场景的步骤和判据由 Agent 按资料与能力决定，面板没有把水文或某套考试步骤写死。复杂业务显示整条依赖链；简单只读查询不强制套大型生产流程。')
h(d,'完整的规划输入','输入包括用户目标与交付范围、真实数据类型/时相/精度、空间参考与单位、工程状态、许可、工具 schema 与实际实现、输出规格，以及适用标准或合同要求。模板只提供候选，不能把关键词匹配结果直接当最优方案。')
h(d,'资料缺口有明确处理方式','任务 intake 记录 requirements、inputs、outputRoot 的缺失/已提供/已确认状态，已确认项引用用户观测。dataNeeds 记录资料用途、推荐规格、影响步骤和核查状态。用户此前明确的信息沿用；缺口只阻断依赖处理，不妨碍独立只读预检。')
h(d,'工具能力选择有清楚顺序','先使用已支持方法；缺工具则搜索受控组合，评估同一科学含义、单位、空间参考、范围、精度和许可下是否等效。替代选择必须写等效条件与限制。确实不能替代时，提出具体开发范围、风险与检查条件；取得用户同意后再开发。不能把 Fill、图像填补、缓冲或其他相似名称混为同一种算法。')
h(d,'行业要求转为可检查判据','将适用条款映射到判据、阈值来源、单位、数据版本、检查范围及必需证据。没有明确适用条款时只做明确说明的通用技术检查，不宣称获得行业认证；软件帮助中的方法说明不等同全部业务规范。')
h(d,'水文例子及能力边界','可先规划DEM预处理、洼地诊断与填洼、流向、汇流、河网阈值与提取、矢量化、出口吸附、流域和指标统计，再安排完整检查。指定出口需核D8与上游范围，真实洼地不一律填平，阈值有单位与依据。目录已含Fill、FlowDirection、Watershed等受控候选，须逐项核schema、许可与真实执行；目录列出不等于完整DEM链已验收，缺失成员先评估有据等效替代。')

page(d,'2  面板与对话怎样同步')
h(d,'一份任务记录，两处用户视图','Agent 将方法、步骤、依赖、观测、检查、问题和成果写入同一原始记录。发布器验证后生成 workflowView，Pro 面板与独立窗口只读展示它；聊天格式化器读取原始记录生成状态卡、流程图与成果列表。两处显示同源，避免各自维护进度。')
h(d,'三个不同的事实层次','业务步骤来自 Agent 的任务记录；请求生命周期来自 MCP 实际接收、执行、返回或异常；受支持 GP 的内部进展来自 SDK 回调。精确 requestId 将请求绑定到某个业务步骤，不根据工具名猜测，以免同名工具在不同任务间串状态。')
h(d,'持续刷新不制造进度','界面每500ms读取变化，Agent 在真实开始/返回/问题/修复/复检/交付时立即发布，并在聊天发送更新。没有新观测显示最后时间；未知总量不造百分比。暂停或关闭窗口只停止观看，不取消 GIS；超时先对账实际效果，不盲重试写入。')
h(d,'执行高亮与查看选择分离','RUNNING 行内容有独立绿色背景、左条、文字标记，选择行只控制右侧详情；可关闭自动跟随查看其他步骤。多个真实执行步骤可同时高亮。没有运行观测则明确显示无当前执行步骤，不把最后一个已完成步骤继续标为执行中。')
picture(d,'native-intake.png','真实 Pro 截图：“资料与交付”与完整业务链同屏；已完成检查与外部验收状态分别展示。')

page(d,'3  检查、修复、复检与最终验收')
h(d,'检查不是工具成功的同义词','工具返回成功首先表示已产出或调用完成；业务检查还要核数据范围、数量、字段、坐标系/单位、参数、方法、覆盖完整性及真实可打开。图件另核真实渲染、图例、比例尺、指北针、标注与用途要求。采样、空数据、不适用与未执行不能冒充全检通过。')
h(d,'质量四阶段与独立验收','预检→过程检查/首检→复检/问题闭环→最终交付检查，各阶段有独立证据。必需阶段未完成或范围不足不能形成完整候选。制作方质量自检形成 PASS CANDIDATE；用户或独立验收方明确裁定另列，执行 Agent 不自签整体 PASS。')
h(d,'失败不可被状态修改抹除','问题记录保留严重性、失败结果、证据、当前修复动作、受影响步骤、下游失效项与新的复检结果。修改数据/方法后更新版本，旧检查失效；替换已完成检查时保留历史。复检必须重新读取并产生新的证据，不能复制首次报告或仅把问题改成已解决。')
h(d,'已实现的几何与拓扑范围','本机对应构件实现 validate_geometries 四项检查和 check_topology_rules 六类明确规则，按数据适用性选择。读取每项结果、总体/扫描数、完整覆盖与限制；这不是 Esri 所有拓扑规则全实现或认证。本机已补充路灯、日照、地面阴影、遮挡率、噪声及真实二维三维综合实验，使用现有MCP与用户明确授权的Pro原生界面组合；这是组合流程证据，不表示缺失算法已新增为MCP工具。其它质量/谱系占位入口和全部PS能力不能据此宣称完成。')
h(d,'交付证据可追查','实际产物逐项列绝对位置、格式、所属步骤和核验状态，计划文件不变成成果下载链接。使用任务版本、资料版本、真实时间、请求响应与哈希核对一致性；哈希证明字节身份，不单独证明科学结果正确。')

page(d,'4  实现结构、可靠性与证据边界')
h(d,'源码职责','ProgressPanelWindow.cs 是原生与独立窗口共用 WPF 视图；WorkflowEventBinding.cs 按准确请求 ID 和观测时间推导有效状态。Skill 的 task_intake.py 校验资料与能力合同，render_task_board.py 生成聊天状态板，publish_workflow.py 校验并发布。安装与启动脚本负责受限路径、哈希、备份和进程环境，不写分析算法。')
h(d,'保护数据和现有工程','发布采用目标 D 目录内临时文件与原子替换；单写锁防止并行发布互相覆盖，任务切换需要准确旧 taskId。损坏、超大或未知版本记录会清除旧的绿色展示并等待有效记录，不复用旧任务成果。受保护工程、历史产物、未知锁与用户编辑默认保留。')
h(d,'安装透明且可恢复','安装前检查包清单哈希，默认 Plan；实际安装需显式选择。遇到不同文件默认拒绝覆盖，明确更新后先备份，逐件记录前后哈希。仅允许本插件 GUID 两处 C 目录变更，其它配套文件在 D。回滚只针对已登记且未被用户后改的文件；不改客户端配置，不公开发布或签名，不自动安装依赖。')
h(d,'本批实际验证','本机 net6/net8 插件与独立窗口均零警告零错误；97项 Skill 辅助检查、16项进度检查通过。真实 Pro 发现239工具，三区12次只读首检/复检通过，15源成员哈希未改；另以300次有界只读调用检查快速事件和运行高亮，已取得真实截帧。300次压力请求不是300种功能覆盖，也不是额外的行业认证。')
h(d,'与普通聊天式操作的明显区别','用户可以先看完整处理和检查链，再跟随当前执行位置；缺数据时知道补什么以及影响哪里；缺工具时能看到受控替代与开发决定；检查失败会留下可追查的修复复检过程；结束时直接定位实际成果。优势来自可核查的一致流程，不用无证据的“全功能自动完成”承诺。')
h(d,'当前仍需如实保留的限制','本批整体独立验收尚未完成，交付为候选。新插件只在 Pro3.5实测；客户端的技能加载和聊天流式能力需分别验证。当前本地发布脚本没有独立的“发布计划”MCP入口。未获得真实内部回调的工具显示未知；通用自动开发、所有行业认证、跨客户端完全同步和所有算法自动替代均不作保证。')
page(d,'5  综合实测带来的科学核验改进')
h(d,'相对高度与地面阴影','开始三维分析前核实际Z与high语义，统一平坦基底只适用于明确的无地形模型。太阳阴影体可能延伸到地面以下，必须核地面交线，不能直接将整个多面体足迹当作地面阴影。')
h(d,'计数与图框的真实含义','轮廓转面会被交叠边界分割，输入建筑数不应硬套为输出分面数。双框分别核Map/Scene类型、四至、朝向、米制比例尺和真实导出；范围缩放可能恢复默认相机朝向，必须写后读回。')
h(d,'自动化与授权边界','缺MCP能力时先评估现有受控替代，用户明确允许后才用原生界面。UI控制先核前台属于Pro，不能将快捷键发给其它软件；无法取得前台时停键鼠输入并保留阻断。原生GP进度和MCP内部事件分别标记。')
h(d,'学习来自证据修正','Skill更新方法选择和核验规则，不会偷偷增加算法、扩白名单或抄旧答案。完整水文工具链仍须另做真实DEM验收，PS链和其它客户端支持另计。')
d.save(OUT/'02_Skill与面板_功能特色与实现原理.docx')
print(str(OUT))
