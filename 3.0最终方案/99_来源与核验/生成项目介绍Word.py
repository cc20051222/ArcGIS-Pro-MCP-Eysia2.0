from pathlib import Path
from datetime import datetime
import re, json, unicodedata, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r'D:\ArcGIS-Pro-MCP 2.0\3.0最终方案')
OUT = ROOT / '06_Word项目介绍'
QA = ROOT / '99_来源与核验' / 'Word排版'
OUT.mkdir(exist_ok=True)
QA.mkdir(exist_ok=True)
OUTPUT = OUT / 'ArcGIS_Pro_MCP_3.0_项目介绍与完整实现方案_最终版_20261002.docx'
DOC = Document()
section = DOC.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.8), Inches(.8)
section.left_margin, section.right_margin = Inches(.8), Inches(.8)
section.header_distance, section.footer_distance = Inches(.35), Inches(.35)

def font_style(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    props = style.element.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); props.append(fonts)
    for attr in ['ascii','hAnsi','eastAsia','cs']:
        fonts.set(qn('w:'+attr), 'Microsoft YaHei')
    for child in list(style.element.iter(qn('w:pBdr'))):
        child.getparent().remove(child)

font_style(DOC.styles['Normal'], 11)
DOC.styles['Normal'].paragraph_format.line_spacing = 1.15
DOC.styles['Normal'].paragraph_format.space_after = Pt(4)
for name,size in [('Title',23),('Subtitle',12),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
    font_style(DOC.styles[name], size, name.startswith('Heading'))
    DOC.styles[name].paragraph_format.space_before = Pt(12 if name.startswith('Heading') else 0)
    DOC.styles[name].paragraph_format.space_after = Pt(7)
    DOC.styles[name].paragraph_format.keep_with_next = True
for name in ['List Bullet','List Number']:
    font_style(DOC.styles[name],11)
    DOC.styles[name].paragraph_format.space_after=Pt(4)

def heading_text(text):
    text = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'[*`]', '', text)
    return re.sub(r'\s+', ' ', ''.join(' ' if unicodedata.category(c)[0] in 'PS' else c for c in text)).strip()

def hyperlink(paragraph, label, href):
    node = OxmlElement('w:hyperlink')
    rid = paragraph.part.relate_to(href, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    node.set(qn('r:id'), rid)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'17456E'); props.append(color)
    underline=OxmlElement('w:u'); underline.set(qn('w:val'),'single'); props.append(underline)
    run.append(props); text=OxmlElement('w:t'); text.text=label; run.append(text); node.append(run); paragraph._p.append(node)

token_re = re.compile(r'(\[([^\]]+)\]\(([^)]*)\)|\*\*([^*]+)\*\*|`([^`]+)`)')
def inline(paragraph, text):
    # The introduction explains arithmetic in prose; exact executable formulae
    # remain in the Markdown and machine model instead of imitated typesetting.
    text=text.replace('K_GP=max(2100,ceil(1.10×N_GP))','为2100与竞品真实同口径GP数量上浮10%后向上取整的较大值')
    text=text.replace('K_sem=max(3000,ceil(1.5×N_sem))','为3000与竞品真实同口径核心语义数量上浮50%后向上取整的较大值')
    text=text.replace('S=w×x1+(1−w)×x2，A=(80,20)、B=(60,60)的差值为60w−40，阈值为2/3',
                      '的评分为第一指标乘权重，再加第二指标乘权重的补数；方案A的两项指标为80和20，方案B为60和60，两者评分差为60乘权重再减40，排序转折权重为三分之二')
    cursor=0
    for m in token_re.finditer(text):
        if m.start()>cursor: paragraph.add_run(text[cursor:m.start()])
        if m.group(2) is not None:
            href=m.group(3).strip('<>')
            if href.startswith(('https://','http://')): hyperlink(paragraph,m.group(2),href)
            else: paragraph.add_run(m.group(2))
        elif m.group(4) is not None:
            paragraph.add_run(m.group(4)).bold=True
        else:
            paragraph.add_run(m.group(5))
        cursor=m.end()
    paragraph.add_run(text[cursor:])

def paragraph(text, style=None):
    p=DOC.add_paragraph(style=style); inline(p,text)
    p.paragraph_format.widow_control=True
    return p

def table(rows):
    n=len(rows[0]); assert n>0
    if n>4:
        # Wide, explanatory matrices become readable records in the introduction.
        for row in rows[1:]:
            p=DOC.add_paragraph(); p.add_run(row[0]).bold=True
            for label,value in zip(rows[0][1:],row[1:]): paragraph(label+'：'+value)
        return
    t=DOC.add_table(rows=0,cols=n); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    widths=([2.05,4.85] if n==2 else [1.15,3.3,2.45] if n==3 else [1.05,1.45,2.2,2.2])
    if rows[0]==['工作项','唯一主责','具体交付']:
        widths=[.8,.65,5.45]
    elif rows[0]==['服务边界','输入与输出','实施与存储责任']:
        widths=[1.85,2.6,2.45]
    elif rows[0]==['对比对象','当前公开优势','我们的采用与超越验收方向']:
        widths=[1.5,2.3,3.1]
    elif rows[0]==['席位','核心职务','主要边界和交付']:
        widths=[.6,1.65,4.65]
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
    t._tbl.tblPr.append(borders)
    for i,values in enumerate(rows):
        cells=t.add_row().cells
        props=t.rows[-1]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit'); props.append(no_split)
        if i==0:
            repeat=OxmlElement('w:tblHeader'); props.append(repeat)
        for j,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                node=OxmlElement('w:'+side); node.set(qn('w:w'),'95'); node.set(qn('w:type'),'dxa'); margin.append(node)
            cp.append(margin)
            fill=OxmlElement('w:shd'); fill.set(qn('w:fill'),'DCE6EF' if i==0 else ('F7F9FA' if i%2==0 else 'FFFFFF')); cp.append(fill)
            p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.1
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(value)<18 else WD_ALIGN_PARAGRAPH.LEFT
            display=value.replace('<br>','\n')
            if j==0 and rows[0][0]=='服务边界':
                display=re.sub(r'(?<=[a-z])(?=[A-Z])',' ',display)
            inline(p,display)
            for run in p.runs: run.font.size=Pt(10.5); run.bold=i==0
    p=DOC.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(0); p.add_run().font.size=Pt(2)

def markdown(text, omit_first_heading=True):
    lines=text.splitlines(); i=0; omitted=False
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('```'):
            kind=line[3:]; block=[]; i+=1
            while i<len(lines) and not lines[i].strip().startswith('```'): block.append(lines[i]); i+=1
            if kind=='mermaid': paragraph('主链执行关系见本介绍的流程图及配套 Markdown 原图。')
            else:
                for code_line in block: paragraph(code_line)
            i+=1; continue
        h=re.match(r'^(#{1,6})\s+(.+)',line)
        if h:
            if len(h.group(1))==1 and omit_first_heading and not omitted: omitted=True; i+=1; continue
            level=min(len(h.group(1)),3)
            DOC.add_heading(heading_text(h.group(2)),level=level)
            i+=1; continue
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:|\-]+\|$', lines[i+1].strip()):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in cells): rows.append(cells)
                i+=1
            maxcols=max(map(len,rows)); rows=[r+['']*(maxcols-len(r)) for r in rows]; table(rows); continue
        if line in ('---','***','___'): i+=1; continue
        line=re.sub(r'^>\s?', '', line)
        bullet=re.match(r'^[-*]\s+(.+)', line)
        number=re.match(r'^\d+[.、)]\s*(.+)', line)
        if bullet: paragraph(bullet.group(1),'List Bullet')
        elif number: paragraph(number.group(1),'List Number')
        else: paragraph(line)
        i+=1

def diagram():
    img=Image.new('RGB',(1800,680),'white'); draw=ImageDraw.Draw(img)
    fontpath=Path(r'C:\Windows\Fonts\msyh.ttc')
    font=ImageFont.truetype(str(fontpath),33); small=ImageFont.truetype(str(fontpath),26)
    boxes=[('任务与输入','目的 来源 版本'),('数据与方法','充分性 适用性'),('计划与准入','类型 批准 预算'),('受控执行','Invoker 作业回执'),('同源事实','指标 图表 说明'),('设计与交付','GIS PS Office')]
    for i,(title,subtitle) in enumerate(boxes):
        x=35+i*292; y=130; draw.rectangle([x,y,x+258,y+165],outline='#666666',width=3)
        draw.text((x+20,y+37),title,font=font,fill='black'); draw.text((x+20,y+95),subtitle,font=small,fill='#333333')
        if i<5:
            ax=x+268; draw.line((ax,y+82,ax+18,y+82),fill='#333333',width=3); draw.polygon([(ax+19,y+82),(ax+11,y+76),(ax+11,y+88)],fill='#333333')
    draw.text((42,355),'共同检查  科学事实  当前权限  真实产物  保存重开  完整版本',font=font,fill='black')
    draw.text((42,430),'更新时沿依赖重新核验   失败时先对账真实效果   成果可预览 可编辑 可重算',font=small,fill='#333333')
    draw.text((42,505),'模型提出候选   确定性合同控制执行   普通用户无需手工操作 Photoshop',font=small,fill='#333333')
    path=QA/'全链路实现关系.png'; img.save(path)
    p=DOC.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Inches(6.85))
    docpr=p._p.xpath('.//wp:docPr')[0]; docpr.set('descr','任务输入经资料方法、计划准入、唯一Invoker执行、同源事实和GIS Photoshop Office交付，所有环节共享检查和更新恢复合同。')

DOC.core_properties.title='ArcGIS Pro MCP 3.0 项目介绍与完整实现方案'
DOC.core_properties.subject='全功能规划 技术实现 社区对比 六席组织与条件工期'
DOC.core_properties.author='ArcGIS Pro MCP 项目规划'
DOC.core_properties.comments='Documentation only; product functionality and full-scope 60-day capacity are not accepted by this document.'
DOC.add_paragraph('ArcGIS Pro MCP 3 0 项目介绍与完整实现方案',style='Title')
DOC.add_paragraph('最终方案  2026年10月3日原位修订',style='Subtitle')
paragraph('本项目面向需要分析、制图和整套成果交付的 GIS 使用者。我们保留原有全部功能和质量要求，建设从资料理解到 GIS 分析、真实 Photoshop 自动设计、Office 输出与后续更新的通用链路。普通用户无需学习 Photoshop；本介绍说明系统如何实现、哪些创新值得采用、怎样验证竞争优势，以及一名指挥端带五名执行端如何推进完整开发。')
paragraph('本资料是最终开发方案。所有新增支持、3000核心语义目标、自动学习及领先指标均需后续真实验收；完整范围60天可行性尚未实测，启动日未确定。安装包仍采用插件及受控本地组件的统一部署形态，完整功能验收后再制作。技术合同、社区对比、工期与五线职责详稿见目录 02–05。')
table([['阅读部分','读者可获得的信息'],['第一部分 产品与创新','完整功能范围 核心思想 十二创新 更新学习和插件发行'],['第二部分 用户与业务链','实际操作步骤 三条标准链 参考成图 数据更新与交接'],['第三部分 技术实现','架构接口 模块实现 科学事实 可靠执行与质量证据'],['第四部分 社区对比','当前公开优势 对比口径 目标如何形成实证'],['第五部分 工期组织与提效','60天条件目标 五执行职责 瓶颈及质量不变的提效']])
diagram()

sources=[]
for section_index,(heading,filename) in enumerate([
 ('第一部分 产品与创新','01_项目最终总方案.md'),
 ('第二部分 用户流程与三条业务链','03_用户操作流程与三条业务链.md'),
 ('第三部分 架构与每个模块的实现','02_架构与模块实施蓝图.md')]):
    p=ROOT/'01_完整最终方案'/filename
    text=p.read_text(encoding='utf-8-sig')
    if section_index < 2: DOC.add_page_break()
    DOC.add_heading(heading,1); markdown(text)
    sources.append({'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest().upper(),'characters':len(text),'selection':'complete body, initial title omitted'})

# Root-authored summaries are checked against the independently researched reports.
supplement=OUT/'项目介绍补充章节.md'
markdown(supplement.read_text(encoding='utf-8-sig'),omit_first_heading=False)
sources.append({'path':str(supplement),'sha256':hashlib.sha256(supplement.read_bytes()).hexdigest().upper(),'selection':'complete body'})

footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run('第 '); run.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
run=footer.add_run(' 页'); run.font.size=Pt(9)
DOC.save(OUTPUT)
report={'file':str(OUTPUT),'bytes':OUTPUT.stat().st_size,'sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest().upper(),
        'title':'ArcGIS Pro MCP 3.0 项目介绍与完整实现方案','sourceBodies':sources,'paragraphs':len(DOC.paragraphs),'tables':len(DOC.tables),
        'nativeLayoutVerification':'PENDING','productVerification':'NOT_RUN','pageSize':'Letter portrait','bodyFontPoints':11}
(QA/'Word生成记录.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
