from __future__ import annotations

import hashlib
import json
import os
import re
import unicodedata
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\ArcGIS-Pro-MCP 2.0\3.0最终方案\宣传")
WORK = ROOT / "制作与核验"
OUT = ROOT / "Word"
COPY_FILE = WORK / "宣传文案.json"
POSTERS = [
    ROOT / "海报" / "01_主宣传海报_竖版.png",
    ROOT / "海报" / "02_产品亮点海报_横版.png",
]
BODY_FONT = "Microsoft YaHei"
BLACK = RGBColor(0, 0, 0)
STATUS = "3.0 方案预告"
BOUNDARY = "发布功能、支持条件和自动化表现以实际验收为准。"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def heading_text(value: str) -> str:
    plain = "".join(" " if unicodedata.category(c).startswith("P") else c for c in value)
    return re.sub(r"\s+", " ", plain).strip()


def east_asia_font(style, name=BODY_FONT):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        if qn(f"w:{key}") in fonts.attrib:
            del fonts.attrib[qn(f"w:{key}")]


def no_border(style):
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        for element in list(ppr):
            if element.tag in {qn("w:pBdr"), qn("w:shd")}:
                ppr.remove(element)


def prepare_document(title: str):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    east_asia_font(normal)
    normal.font.size = Pt(11.3)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.14
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    sizes = {"Title": 23, "Subtitle": 12, "Heading 1": 21, "Heading 2": 13.5, "Heading 3": 12}
    for name, size in sizes.items():
        style = doc.styles[name]
        east_asia_font(style)
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.underline = False
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(8 if name == "Heading 2" else 0)
        style.paragraph_format.space_after = Pt(7 if name != "Heading 1" else 10)
        style.paragraph_format.keep_with_next = True
        no_border(style)

    for name in ("List Bullet", "List Paragraph", "Header", "Footer"):
        style = doc.styles[name]
        east_asia_font(style)
        style.font.color.rgb = BLACK
        style.font.size = Pt(11.3 if name.startswith("List") else 9)
        no_border(style)
    bullets = doc.styles["List Bullet"]
    bullets.paragraph_format.space_after = Pt(4)
    bullets.paragraph_format.line_spacing = 1.12

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    footer.paragraph_format.line_spacing = 1.0
    run = footer.add_run("3.0 方案预告  发布功能及支持条件以实际验收为准  ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = BLACK
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    doc.core_properties.title = title
    doc.core_properties.subject = "ArcGIS Pro MCP 3.0 规划功能与目标体验宣传资料"
    doc.core_properties.author = "ArcGIS Pro MCP 项目"
    doc.core_properties.keywords = "GIS,Photoshop,Office,3.0方案预告,宣传"
    doc.core_properties.comments = "概念宣传资料，不代表当前软件截图或已实现能力。"
    return doc


def add_status(doc, label=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{STATUS}  {label}".strip())
    r.font.size = Pt(10)
    r.font.color.rgb = BLACK
    return p


def add_intro(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(9)
    return p


def add_poster(doc, poster: Path, index: int, standalone: bool):
    if standalone and index == 0:
        doc.add_paragraph("ArcGIS Pro MCP 3 0 宣传海报", style="Title")
    else:
        heading = "主宣传海报" if index == 0 else "产品亮点海报"
        doc.add_paragraph(heading_text(heading), style="Heading 1")
    add_status(doc, "概念宣传视觉")

    with Image.open(poster) as im:
        width, height = im.size
    max_width = 6.7
    max_height = 7.7 if width < height else 5.3
    scale = min(max_width / width, max_height / height)
    fitted_width = width * scale
    fitted_height = height * scale

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3 if width < height else 34)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    picture = p.add_run().add_picture(str(poster), width=Inches(fitted_width), height=Inches(fitted_height))
    picture._inline.docPr.set("descr", "ArcGIS Pro MCP 3.0 概念宣传海报；图像与布局为宣传视觉，不是实际软件截图或真实业务成果。")

    note = doc.add_paragraph("规划功能与目标体验。" + BOUNDARY)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(0)
    for r in note.runs:
        r.font.size = Pt(9)
    return {"source": str(poster), "sourcePixels": [width, height], "wordInches": [fitted_width, fitted_height]}


def write_brochure(copy):
    doc = prepare_document("ArcGIS Pro MCP 3 0 产品宣传册")
    for i, page in enumerate(copy["pages"]):
        if i:
            doc.add_page_break()
        if i == 0:
            doc.add_paragraph("ArcGIS Pro MCP 3 0 产品宣传册", style="Title")
        add_status(doc, page["eyebrow"])
        doc.add_paragraph(heading_text(page["title"]), style="Heading 1")
        add_intro(doc, page["introduction"])
        for group in page["sections"]:
            doc.add_paragraph(heading_text(group["heading"]), style="Heading 2")
            if group.get("bullets"):
                for bullet in group["bullets"]:
                    doc.add_paragraph(bullet, style="List Bullet")
            if group.get("paragraph"):
                doc.add_paragraph(group["paragraph"])
    placement = []
    for i, poster in enumerate(POSTERS):
        doc.add_page_break()
        placement.append(add_poster(doc, poster, i, False))
    path = OUT / "ArcGIS_Pro_MCP_3.0_产品宣传册_预告版.docx"
    doc.save(path)
    return path, placement


def write_posters():
    doc = prepare_document("ArcGIS Pro MCP 3 0 宣传海报")
    placement = []
    for i, poster in enumerate(POSTERS):
        if i:
            doc.add_page_break()
        placement.append(add_poster(doc, poster, i, True))
    path = OUT / "ArcGIS_Pro_MCP_3.0_宣传海报_展示版.docx"
    doc.save(path)
    return path, placement


def check_structure(path, expected_breaks):
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        images = [n for n in archive.namelist() if n.startswith("word/media/")]
        breaks = document_xml.count('w:type="page"')
        if breaks != expected_breaks or len(images) != 2:
            raise RuntimeError(f"Unexpected breaks or images in {path.name}: {breaks}, {len(images)}")
        if "w:txbxContent" in document_xml:
            raise RuntimeError("Unapproved text box found")
    return {"manualPageBreaks": breaks, "embeddedImages": len(images), "tables": len(Document(path).tables)}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    os.environ["TEMP"] = str(WORK)
    os.environ["TMP"] = str(WORK)
    os.environ["TMPDIR"] = str(WORK)
    missing = [str(path) for path in POSTERS if not path.is_file()]
    if missing:
        raise SystemExit(json.dumps({"status": "WAITING_FOR_POSTER_ASSETS", "missing": missing}, ensure_ascii=False))
    copy = json.loads(COPY_FILE.read_text(encoding="utf-8-sig"))
    if len(copy["pages"]) != 8:
        raise RuntimeError("Expected eight brochure pages")

    brochure, brochure_placement = write_brochure(copy)
    poster_doc, poster_placement = write_posters()
    receipt = {
        "status": "AUTHORING_COMPLETE_PENDING_NATIVE_RENDER",
        "sourceCopy": {"path": str(COPY_FILE), "sha256": digest(COPY_FILE)},
        "sourcePosters": [{"path": str(p), "sha256": digest(p)} for p in POSTERS],
        "outputs": [
            {"path": str(brochure), "sha256": digest(brochure), "bytes": brochure.stat().st_size, "plannedPages": 10, "structuralChecks": check_structure(brochure, 9), "posterPlacement": brochure_placement},
            {"path": str(poster_doc), "sha256": digest(poster_doc), "bytes": poster_doc.stat().st_size, "plannedPages": 2, "structuralChecks": check_structure(poster_doc, 1), "posterPlacement": poster_placement},
        ],
        "renderStatus": "NOT_RUN_BY_AUTHORING_AGENT",
        "notes": ["根代理负责原生渲染和实际页面复核。", "手工分页与OOXML检查不能替代真实Word分页和视觉核验。"],
    }
    (WORK / "宣传Word制作回执.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
